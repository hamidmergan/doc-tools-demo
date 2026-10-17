import json
import logging

import docx
import gradio as gr
import pytest
from doc_chunker import NoTextError
from doc_extract import ConfigError

from doc_tools_demo import app, core


# Settings


def test_settings_from_env(monkeypatch):
    monkeypatch.setenv("DOC_TOOLS_DEMO_BASE_URL", "http://localhost:8000/v1")
    monkeypatch.setenv("DOC_TOOLS_DEMO_MODEL", "text")
    monkeypatch.setenv("DOC_TOOLS_DEMO_VISION_MODEL", " ")
    s = core.Settings.from_env()
    assert s.base_url == "http://localhost:8000/v1"
    assert s.model == "text"
    assert s.vision_model is None
    assert s.api_key is None


def test_settings_blank_values_become_none():
    s = core.Settings(" http://x/v1 ", "m", "", "  ")
    assert s.text_args() == {"base_url": "http://x/v1", "model": "m", "api_key": None}
    assert s.vision_args()["model"] == "m"
    assert core.Settings("u", "m", "v").vision_args()["model"] == "v"


# Glossary


def test_build_glossary_rows_win_over_file(tmp_path):
    file = tmp_path / "terms.csv"
    file.write_text("source,target\npump,Pumpe\nvalve,Ventil\n", encoding="utf-8")
    rows = [["valve", "Absperrventil"], ["", "ignored"], ["API", ""], [None, None]]
    assert core.build_glossary(rows, str(file)) == {"pump": "Pumpe", "valve": "Absperrventil", "API": ""}


def test_build_glossary_empty():
    assert core.build_glossary([["", ""]], None) is None
    assert core.build_glossary(None, None) is None


# Pages


def test_split_pages():
    md = "<!-- page 1 -->\nFirst\n\n<!-- page 2 -->\nSecond\n"
    assert [(p, t.strip()) for p, t in core.split_pages(md)] == [(1, "First"), (2, "Second")]
    assert core.split_pages("Just text") == [(None, "Just text")]


# Chunk


def test_chunk_file_markdown(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("# Notes\n\n" + "A sentence about pumps. " * 40, encoding="utf-8")
    chunks = core.chunk_file(str(path), 300, 30)
    assert len(chunks) > 1
    assert all(len(c.text) <= 400 for c in chunks)


def test_to_jsonl():
    text = core.to_jsonl([{"a": "ü"}, {"b": 2}])
    assert [json.loads(line) for line in text.splitlines()] == [{"a": "ü"}, {"b": 2}]


# Extract


def test_extract_uses_vision_model(server, settings, make_image):
    md = core.extract_markdown(make_image(), settings)
    assert "Pump manual" in md
    assert server.models() == ["vision-model"]


def test_extract_needs_settings(make_image):
    with pytest.raises(ConfigError):
        core.extract_markdown(make_image(), core.Settings())


# Translate


def test_translate_plain(server, settings):
    translator = core.make_translator(settings, "German")
    text, warnings = core.translate_plain("first line\n\nsecond line", translator)
    assert text == "FIRST LINE\n\nSECOND LINE"
    assert warnings == []
    assert set(server.models()) == {"text-model"}


def test_translate_markdown_file(server, settings, tmp_path):
    src = tmp_path / "readme.md"
    src.write_text("# Title\n\nSome text.\n", encoding="utf-8")
    translator = core.make_translator(settings, "de")
    out, _ = core.translate_document(str(src), core.work_dir(), translator)
    assert out.exists()
    assert out.parent.parent == core.work_root()
    content = out.read_text(encoding="utf-8")
    assert "TITLE" in content and "SOME TEXT." in content


def test_translate_docx_file(server, settings, tmp_path):
    src = tmp_path / "letter.docx"
    doc = docx.Document()
    doc.add_paragraph("Hello world")
    doc.save(src)
    translator = core.make_translator(settings, "de")
    out, _ = core.translate_document(str(src), core.work_dir(), translator)
    assert out.suffix == ".docx"
    assert [p.text for p in docx.Document(out).paragraphs] == ["HELLO WORLD"]


def test_glossary_warning_is_collected(server, settings):
    translator = core.make_translator(settings, "German", glossary={"pump": "Kreiselpumpe"})
    _, warnings = core.translate_plain("The pump is loud.", translator)
    assert any("Kreiselpumpe" in w for w in warnings)


def test_collect_warnings_only_keeps_this_thread():
    import threading

    log = logging.getLogger("office_translate.translator")
    with core.collect_warnings() as found:
        log.warning("mine")
        t = threading.Thread(target=log.warning, args=("other",))
        t.start()
        t.join()
    log.warning("after")
    assert found == ["mine"]


# Pipeline


def test_pipeline_without_translation(server, settings, make_pdf):
    result = core.run_pipeline(make_pdf(pages=2), settings, chunk_size=200, overlap=0)
    assert "<!-- page 2 -->" in result.markdown
    assert [r["page"] for r in result.rows] == [1, 2]
    assert all("translation" not in r for r in result.rows)
    assert result.markdown_file.read_text(encoding="utf-8") == result.markdown
    lines = result.jsonl_file.read_text(encoding="utf-8").splitlines()
    assert json.loads(lines[0])["source"] == "scan.pdf"


def test_pipeline_with_translation(server, settings, make_image):
    translator = core.make_translator(settings, "German")
    result = core.run_pipeline(make_image(), settings, translator)
    assert result.rows
    row = result.rows[0]
    assert row["page"] is None
    assert row["translation"] == row["text"].upper()
    assert row["target_lang"] == "German"
    assert "vision-model" in server.models() and "text-model" in server.models()


def test_pipeline_no_text(settings, make_image, monkeypatch):
    monkeypatch.setattr(core, "extract_markdown", lambda path, s: "   ")
    with pytest.raises(NoTextError):
        core.run_pipeline(make_image(), settings)


# App


def test_app_builds():
    assert isinstance(app.build(core.Settings()), gr.Blocks)


def test_cli_version(capsys):
    with pytest.raises(SystemExit):
        app.main(["--version"])
    assert "doc-tools-demo" in capsys.readouterr().out


def test_handlers_ask_for_input():
    with pytest.raises(gr.Error):
        app.run_extract(None, "", "", "", "")
    with pytest.raises(gr.Error):
        app.run_chunk(None, 1000, 150)
    with pytest.raises(gr.Error):
        app.run_translate(None, " ", "German", "", "", None, None, False, "", "", "", "")


def test_handler_explains_missing_settings(make_image):
    with pytest.raises(gr.Error) as info:
        app.run_extract(str(make_image()), "", "", "", "")
    assert "Model settings" in info.value.message


def test_handler_translate_text(server):
    out_file, text, notes = app.run_translate(
        None, "good morning", "German", "", "", [["", ""]], None, False,
        server.base_url, "text-model", "", "",
    )
    assert out_file is None
    assert text == "GOOD MORNING"
    assert notes == ""


def test_handler_pipeline_ignores_glossary_when_unticked(server, make_image):
    md, table, files, info = app.run_pipeline(
        str(make_image()), "German", 1000, 150, False,
        "English", "be formal", [["pump", "Pumpe"]], None, True,
        server.base_url, "text-model", "vision-model", "",
    )
    assert "Pump manual" in md
    assert table[0][3] == table[0][2].upper()
    assert len(files) == 2
    assert "Warnings" not in info
